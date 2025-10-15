from __future__ import annotations
import os, json, hashlib
from pathlib import Path
from typing import Any, Dict, List, Tuple, Sequence, Optional

import tiktoken

# --- optional imports by file type ---
try:
    import fitz  # PyMuPDF for PDFs
except Exception:
    fitz = None  # graceful fallback

try:
    from docx import Document  # python-docx for DOCX
except Exception:
    Document = None

try:
    import openpyxl  # XLSX
except Exception:
    openpyxl = None

try:
    import csv
except Exception:
    csv = None

try:
    from PIL import Image
    import pytesseract
except Exception:
    Image = None
    pytesseract = None

from services.vector_langchain import add_texts  # LangChain+Chroma wrapper

# ---------------- paths ----------------
def _store_root() -> Path:
    return Path(os.getenv("VECTOR_STORE_PATH", "./src/data/indexes")).resolve()

def _chunks_dir(framework_dir: Path) -> Path:
    d = framework_dir / "chunks"
    d.mkdir(parents=True, exist_ok=True)
    return d

# ---------------- helpers ----------------
def _sha256(s: str) -> str:
    return hashlib.sha256(s.encode("utf-8")).hexdigest()

def _enc():
    return tiktoken.get_encoding("cl100k_base")

def _chunk_by_tokens(text: str, chunk_size: int = 500, overlap: int = 80) -> List[str]:
    enc = _enc()
    toks = enc.encode(text or "")
    out: List[str] = []
    start = 0
    while start < len(toks):
        end = min(start + chunk_size, len(toks))
        piece = enc.decode(toks[start:end])
        out.append(piece)
        if end >= len(toks):
            break
        start = max(0, end - overlap)
    return out

def _norm_ws(s: str) -> str:
    return " ".join((s or "").split())

# ---------------- extractors ----------------
def _extract_pdf(path: Path) -> List[Tuple[int, str]]:
    if not fitz:
        return [(1, f"[PDF parsing unavailable] {path.name}")]
    pages: List[Tuple[int, str]] = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc):
            txt = _norm_ws(page.get_text("text") or "")
            if txt:
                pages.append((i+1, txt))
    return pages or [(1, "")]

def _extract_txt(path: Path) -> List[Tuple[int, str]]:
    try:
        txt = _norm_ws(path.read_text(encoding="utf-8"))
    except Exception:
        txt = _norm_ws(path.read_text(errors="ignore"))
    return [(1, txt)]

def _extract_docx(path: Path) -> List[Tuple[int, str]]:
    if Document is None:
        return [(1, f"[DOCX parsing unavailable] {path.name}")]
    doc = Document(str(path))
    paras = [p.text for p in doc.paragraphs]
    txt = _norm_ws("\n".join(paras))
    return [(1, txt)]

def _extract_xlsx(path: Path) -> List[Tuple[int, str]]:
    if openpyxl is None:
        return [(1, f"[XLSX parsing unavailable] {path.name}")]
    wb = openpyxl.load_workbook(str(path), data_only=True)
    parts: List[Tuple[int, str]] = []
    idx = 1
    for ws in wb.worksheets:
        lines: List[str] = [f"[Sheet] {ws.title}"]
        for row in ws.iter_rows(values_only=True):
            row_str = " | ".join("" if v is None else str(v) for v in row)
            lines.append(row_str)
        chunk = _norm_ws("\n".join(lines))
        if chunk:
            parts.append((idx, chunk))
            idx += 1
    return parts or [(1, "")]

def _extract_csv(path: Path) -> List[Tuple[int, str]]:
    if csv is None:
        return [(1, f"[CSV parsing unavailable] {path.name}")]
    lines: List[str] = []
    with path.open("r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        for row in reader:
            lines.append(" | ".join(row))
    txt = _norm_ws("\n".join(lines))
    return [(1, txt or "")]
    
def _extract_image(path: Path) -> List[Tuple[int, str]]:
    if Image is None or pytesseract is None:
        return [(1, f"[OCR unavailable] {path.name}")]
    try:
        img = Image.open(path)
        txt = pytesseract.image_to_string(img) or ""
        return [(1, _norm_ws(txt))]
    except Exception:
        return [(1, "")]

def _extract_generic(path: Path) -> List[Tuple[int, str]]:
    # last resort: try to read as text
    try:
        txt = path.read_text(encoding="utf-8")
    except Exception:
        txt = path.read_text(errors="ignore")
    return [(1, _norm_ws(txt))]

def _extract_by_ext(path: Path) -> List[Tuple[int, str]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".txt":
        return _extract_txt(path)
    if ext in (".docx",):
        return _extract_docx(path)
    if ext in (".xlsx",):
        return _extract_xlsx(path)
    if ext in (".csv",):
        return _extract_csv(path)
    if ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff"):
        return _extract_image(path)
    # fallback
    return _extract_generic(path)

# ---------------- public indexers ----------------
def index_framework(framework: str, src_root: Path) -> Dict[str, Any]:
    """
    Reads chunked JSONL from: src_root/guidelines/<framework>/chunks/chunks.jsonl
    Loads into vector store collection: fw_<framework>
    """
    chunks_file = src_root / "guidelines" / framework / "chunks" / "chunks.jsonl"
    if not chunks_file.exists():
        raise FileNotFoundError(f"No chunks found at {chunks_file}")
    texts, metas, ids = [], [], []
    with chunks_file.open("r", encoding="utf-8") as f:
        for i, line in enumerate(f, start=1):
            obj = json.loads(line)
            text = obj.get("text") or ""
            if not text.strip():
                continue
            meta = {
                "framework": obj.get("framework"),
                "source_pdf": obj.get("source_pdf"),
                "page": obj.get("page"),
                "chunk_index": obj.get("chunk_index"),
                "sha256": obj.get("sha256"),
            }
            metas.append(meta)
            texts.append(text)
            ids.append(obj.get("sha256") or f"{i}")
    add_texts(collection_name=f"fw_{framework}", ids=ids, texts=texts, metadatas=metas)
    return {"framework": framework, "count": len(texts)}

def index_assessment_pdf(firm: str, pdf_path: Path) -> Dict[str, Any]:
    parts = _extract_pdf(pdf_path)
    return _index_evidence_like(
        collection=f"assessment_{firm}",
        doc_id=pdf_path.name,
        parts=parts,
        source_type="assessment_pdf",
        ext="pdf"
    )

def index_evidence_file(firm: str, path: Path) -> Dict[str, Any]:
    parts = _extract_by_ext(path)
    return _index_evidence_like(
        collection=f"evidence_{firm}",
        doc_id=path.name,
        parts=parts,
        source_type="evidence",
        ext=path.suffix.lower().lstrip(".") or "bin"
    )

def index_evidence_batch(firm: str, paths: List[Path]) -> Dict[str, Any]:
    total_chunks = 0
    files_out: List[Dict[str, Any]] = []
    for p in paths:
        try:
            info = index_evidence_file(firm, p)
            total_chunks += info.get("count", 0)
            files_out.append({"file": p.name, "chunks": info.get("count", 0), "status": "ok"})
        except Exception as e:
            files_out.append({"file": p.name, "chunks": 0, "status": "error", "error": str(e)})
    return {"total_docs": len(paths), "total_chunks": total_chunks, "files": files_out}

# ---------------- internal ----------------
def _index_evidence_like(
    *,
    collection: str,
    doc_id: str,
    parts: List[Tuple[int, str]],  # [(page_or_part, text)]
    source_type: str,
    ext: str,
) -> Dict[str, Any]:
    ids: List[str] = []
    texts: List[str] = []
    metas: List[Dict[str, Any]] = []
    count = 0
    for page, raw in parts:
        if not raw.strip():
            continue
        # chunk each part
        chunks = _chunk_by_tokens(raw, chunk_size=500, overlap=80)
        for ci, ch in enumerate(chunks, start=1):
            sha = _sha256(f"{doc_id}:{page}:{ci}:{ch[:64]}")
            ids.append(sha)
            texts.append(ch)
            metas.append({
                "doc_id": doc_id,
                "page": page,
                "chunk_index": ci,
                "source_type": source_type,
                "ext": ext,
            })
            count += 1
    if count:
        add_texts(collection_name=collection, ids=ids, texts=texts, metadatas=metas)
    return {"doc_id": doc_id, "count": count}
